from docx import Document
from docx.shared import Inches

# Create document
doc = Document()
doc.add_heading("Test Document", 0)
doc.add_paragraph("This is a test paragraph for file2ai conversion testing.")
doc.add_paragraph("It includes multiple paragraphs to test text extraction.")

# Add a table
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Header 1"
table.cell(0, 1).text = "Header 2"
table.cell(1, 0).text = "Value 1"
table.cell(1, 1).text = "Value 2"

doc.save("sample.docx")
print("Test document created successfully!")
